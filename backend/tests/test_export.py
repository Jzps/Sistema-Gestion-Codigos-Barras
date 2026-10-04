"""Exportacion XLSX del historial."""

from io import BytesIO

from openpyxl import load_workbook

from app.services.excel_export import HEADERS
from tests.conftest import login, make_user


def _setup_with_scans(client, db):
    make_user(db, username="op", password="secret123", workspace_name="WS Export")
    login(client, "op", "secret123")
    client.post(
        "/api/scans",
        json={
            "barcode": "EXP-001",
            "weight_value": "44.09",
            "weight_unit": "LB",
            "product_name": "Caja export",
        },
    )
    client.post("/api/scans", json={"barcode": "EXP-001"})


def test_export_returns_valid_xlsx_with_history(client, db):
    _setup_with_scans(client, db)
    response = client.get("/api/reports/scans.xlsx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == HEADERS
    assert ws.max_row == 3  # encabezado + 2 escaneos

    data = list(ws.iter_rows(min_row=2, values_only=True))
    assert all(row[2] == "EXP-001" for row in data)  # columna Código
    assert all(row[7] == "op" for row in data)  # columna Usuario
    assert data[0][4] == 44.09  # Peso
    assert data[0][5] == "LB"  # Unidad
    assert abs(data[0][6] - 20) < 0.01  # Peso (kg)
    assert ws.freeze_panes == "A2"


def test_export_requires_auth(client):
    assert client.get("/api/reports/scans.xlsx").status_code == 401
