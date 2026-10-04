"""Exportacion del historial de escaneos a .xlsx con openpyxl.

El archivo se genera en memoria bajo demanda; no se persiste en disco.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.schemas.scan import ScanOut

HEADERS = ["Fecha", "Hora", "Código", "Producto", "Peso", "Unidad", "Peso (kg)", "Usuario"]

_COLUMN_WIDTHS = [12, 10, 20, 28, 10, 8, 12, 16]


def build_scans_xlsx(rows: list[ScanOut]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Historial"

    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(
            [
                row.scanned_at.strftime("%Y-%m-%d"),
                row.scanned_at.strftime("%H:%M:%S"),
                row.barcode_raw,
                row.product_name or "",
                float(row.weight_value),
                row.weight_unit,
                float(row.weight_kg),
                row.username,
            ]
        )

    # Usabilidad: encabezado fijo, autofiltro y anchos razonables.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
